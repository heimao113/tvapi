#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
播放链路实测:对源清单中"存活"的源走一遍 取样片→详情→m3u8→TS切片(前256KB),
测出真实播放速度,产出 output/测速排行.csv 并写入 源清单.xlsx 的"播放测速"sheet。
"""
import csv
import ipaddress
import json
import os
import socket
import time
import urllib.parse
import concurrent.futures as cf
import urllib3
import requests

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "output")
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
WORKERS = 16
SAMPLE_BYTES = 256 * 1024

CF_RANGES = [ipaddress.ip_network(n) for n in (
    "104.16.0.0/13", "104.24.0.0/14", "172.64.0.0/13", "131.0.72.0/22",
    "141.101.0.0/16", "162.158.0.0/15", "188.114.0.0/17", "198.41.128.0/17",
)]

session = requests.Session()
session.headers.update({"User-Agent": UA})
session.verify = False


def is_cf(host: str) -> bool:
    try:
        ip = ipaddress.ip_address(socket.gethostbyname(host.split(":")[0]))
    except Exception:
        return False
    return any(ip in net for net in CF_RANGES)


def get_json(api: str, params: str):
    r = session.get(api + params, timeout=12)
    return r.json()


def first_m3u8(api: str) -> tuple:
    j = get_json(api, "?ac=videolist&pg=1")
    vids = j.get("list") or []
    if not vids:
        return None, None
    title = vids[0].get("vod_name", "")
    play = vids[0].get("vod_play_url") or ""
    if not play:  # 个别站列表不带播放地址,补一次 ids 详情
        det = get_json(api, f"?ac=videolist&ids={vids[0].get('vod_id')}").get("list") or []
        play = det[0].get("vod_play_url", "") if det else ""
    for seg in play.split("$$$")[0].split("#"):
        url = seg.split("$")[-1]
        if ".m3u8" in url and url.startswith("http"):
            return url, title
    return None, title


def test_one(row: dict) -> dict:
    api, name = row["接口"].rstrip("/"), row["名称"]
    res = {"名称": name, "接口": api, "是否成人": row["是否成人"],
           "样例片名": "", "m3u8延迟ms": "-", "TS速度KB/s": "-", "评级": "未测出", "Cloudflare CDN": "-"}
    try:
        res["Cloudflare CDN"] = "是" if is_cf(urllib.parse.urlsplit(api).netloc) else "否"
        t0 = time.time()
        m3u8, title = first_m3u8(api)
        res["样例片名"] = (title or "")[:24]
        if not m3u8:
            return res
        res["m3u8延迟ms"] = int((time.time() - t0) * 1000)
        r = session.get(m3u8, timeout=(5, 8))
        body = r.text
        # 主播放列表 → 取第一个子 m3u8
        if "#EXT-X-STREAM-INF" in body:
            for line in body.splitlines():
                line = line.strip()
                if line and not line.startswith("#"):
                    m3u8 = urllib.parse.urljoin(m3u8, line)
                    body = session.get(m3u8, timeout=(5, 8)).text
                    break
        ts = next((urllib.parse.urljoin(m3u8, l.strip())
                   for l in body.splitlines() if l.strip() and not l.strip().startswith("#")), None)
        if not ts:
            return res
        t1 = time.time()
        got = 0
        with session.get(ts, timeout=(5, 15), stream=True) as tr:
            for chunk in tr.iter_content(32 * 1024):
                got += len(chunk)
                if got >= SAMPLE_BYTES:
                    break
        dt = time.time() - t1
        if got < 64 * 1024:
            res["评级"] = "样例过小"
            return res
        kbps = int(got / dt / 1024)
        res["TS速度KB/s"] = kbps
        res["评级"] = "快" if kbps >= 1000 else ("中" if kbps >= 250 else "慢")
    except Exception as e:
        res["评级"] = f"失败:{type(e).__name__}"
    return res


def main():
    with open(os.path.join(OUT, "源清单.csv"), encoding="utf-8-sig") as f:
        rows = [r for r in csv.DictReader(f) if r["状态"] == "存活"]
    print(f"[测速] 待测 {len(rows)} 个存活源 …")
    results = []
    with cf.ThreadPoolExecutor(max_workers=WORKERS) as ex:
        futs = [ex.submit(test_one, r) for r in rows]
        for i, fu in enumerate(cf.as_completed(futs), 1):
            results.append(fu.result())
            if i % 20 == 0:
                print(f"[测速] {i}/{len(rows)}")

    def spd(r):
        v = r["TS速度KB/s"]
        return v if isinstance(v, int) else -1
    results.sort(key=lambda r: (r["是否成人"], -spd(r)))

    cols = list(results[0].keys())
    with open(os.path.join(OUT, "测速排行.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(results)

    try:
        from openpyxl import load_workbook
        path = os.path.join(OUT, "源清单.xlsx")
        wb = load_workbook(path)
        if "播放测速" in wb.sheetnames:
            del wb["播放测速"]
        ws = wb.create_sheet("播放测速")
        ws.append(cols)
        for r in results:
            ws.append([r[c] for c in cols])
        wb.save(path)
    except Exception as e:
        print(f"[警告] Excel 写入失败: {e}", file=sys.stderr)

    fast = [r for r in results if r["是否成人"] == "否" and isinstance(r["TS速度KB/s"], int)]
    print(f"[测速完成] 有效 {len([r for r in results if isinstance(r['TS速度KB/s'], int)])}/{len(rows)}")
    print("人版最快的 10 个源:")
    for r in fast[:10]:
        print(f"  {r['TS速度KB/s']:>6} KB/s | {r['名称']} | m3u8 {r['m3u8延迟ms']}ms")


if __name__ == "__main__":
    import sys
    sys.stderr = sys.stderr  # keep stderr
    main()
